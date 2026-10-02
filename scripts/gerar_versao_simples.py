"""Gera a versão simples (para leigos) a partir do relatório completo já recalculado.

Uso: python scripts/gerar_versao_simples.py [relatorio_completo.xlsx] [saida.xlsx]
Os números são lidos do relatório completo, então batem exatamente com ele.
"""
import sys
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

RAIZ = Path(__file__).resolve().parent.parent
SRC = sys.argv[1] if len(sys.argv) > 1 else str(RAIZ / 'planilhas' / 'ZBR_Importado_Relatorio_Gestao.xlsx')
OUT = sys.argv[2] if len(sys.argv) > 2 else str(RAIZ / 'planilhas' / 'ZBR_Importado_Versao_Simples.xlsx')

F = 'Calibri'
NAVY, CORP, LIGHT, ZEBRA, DARK, GREY = '17365D', '245A81', 'D9EAF7', 'F2F4F7', '404854', '6B7480'
G_BG, G_TX, Y_BG, Y_TX, R_BG, R_TX = 'E2EFDA', '375623', 'FFF2CC', '7F6000', 'F8DADA', 'C00000'
NUM = '#,##0'
SIGN = '+#,##0;-#,##0;0'

src = openpyxl.load_workbook(SRC, data_only=True)
conc = src['5 Conciliação']
HR = next(r for r in range(1, 30) if conc.cell(row=r, column=1).value == 'SKU')

# ---------- linguagem simples ----------
ACONTECEU = {
    'Falta sem explicação': 'Tem menos do que deveria e não achamos o motivo nos dados.',
    'Não localizado no estoque': 'Entrou no sistema, mas não foi encontrado no estoque.',
    'Saída sem baixa no estoque': 'Foi vendido ou entregue, mas ainda aparece no estoque (faltou dar baixa).',
    'Doca/transferência pendente': 'Está parado na doca ou em transferência, ainda não foi guardado.',
    'Estoque anterior a 09/09': 'É estoque antigo, de antes de 9 de setembro.',
    'Troca de lote': 'Foi lançado no lote errado: a mesma quantidade aparece em outro lote.',
    'Sobra parcialmente explicada': 'Tem mais do que deveria; parte é estoque antigo, o resto precisa ser conferido.',
    'Sobra sem explicação': 'Tem mais do que deveria e não achamos o motivo nos dados.',
    '—': 'Tudo certo.'}
FAZER = {
    'Falta sem explicação': 'Contar o lote. Se faltar mesmo, avisar para corrigir no sistema.',
    'Não localizado no estoque': 'Procurar no armazém. Se não achar, contar e corrigir no sistema.',
    'Saída sem baixa no estoque': 'Pedir para dar baixa das entregas no sistema.',
    'Doca/transferência pendente': 'Guardar a mercadoria na posição ou dar baixa da entrega.',
    'Estoque anterior a 09/09': 'Confirmar o estoque de 8 de setembro. Se estiver certo, nada a fazer.',
    'Troca de lote': 'Corrigir o lote no sistema.',
    'Sobra parcialmente explicada': 'Contar o lote e corrigir a diferença no sistema.',
    'Sobra sem explicação': 'Contar o lote e corrigir a diferença no sistema.',
    '—': '—'}
PRIOR = {'Crítica': 'Urgente', 'Alta': 'Alta', 'Média': 'Média', 'Baixa': 'Baixa', '—': ''}
SIT = {'Confere': '✔ Certo', 'Sobrando': '▲ Sobrando', 'Faltando': '▼ Faltando'}
ORD = {'Urgente': 0, 'Alta': 1, 'Média': 2, 'Baixa': 3, '': 4}

linhas = []
for r in range(HR + 1, conc.max_row + 1):
    sku = conc.cell(row=r, column=1).value
    if not sku or not conc.cell(row=r, column=2).value or conc.cell(row=r, column=10).value not in SIT:
        continue  # ignora notas de rodapé
    v = lambda c: conc.cell(row=r, column=c).value  # noqa: E731
    cat, crit = v(17) or '—', v(18) or '—'
    linhas.append(dict(sku=sku, lote=v(2), desc=v(3) or '', um=(v(4) or '').lower(), deveria=v(7) or 0,
                       tem=v(8) or 0, dif=v(9) or 0, sit=SIT.get(v(10), v(10)), cat=cat,
                       aconteceu=ACONTECEU.get(cat, cat), fazer=FAZER.get(cat, ''), prior=PRIOR.get(crit, crit)))
linhas.sort(key=lambda d: (ORD.get(d['prior'], 9), -abs(d['dif']), d['sku'], d['lote']))
problemas = [d for d in linhas if d['sit'] != '✔ Certo']

# depósito 2008: mercadoria parada (a partir da base de entrada + saldo inicial)
ent = src['3 Base Entrada 2008']
g = {}
for r in range(2, ent.max_row + 1):
    sku, lote = ent.cell(row=r, column=2).value, ent.cell(row=r, column=3).value
    if not sku:
        continue
    q, um, mov = ent.cell(row=r, column=7).value or 0, ent.cell(row=r, column=8).value, ent.cell(row=r, column=11).value
    o = g.setdefault((sku, lote), dict(um=(um or '').lower(), chegou=0, saiu=0, ini=0))
    if mov == 'Recebido no 2008':
        o['chegou'] += q
    elif mov == 'Entrada':
        o['saiu'] += -q
    else:
        o['saiu'] -= q  # estorno: voltou para o 2008
si = src['3 Base Saldo Inicial 08-09']
for r in range(2, si.max_row + 1):
    sku, lote, dep, q = (si.cell(row=r, column=c).value for c in (1, 2, 3, 5))
    if sku and str(dep) == '2008' and (sku, lote) in g:
        g[(sku, lote)]['ini'] += q or 0
desc = {d['sku']: d['desc'] for d in linhas if d['desc']}
parado = sorted(((k, o, round(o['ini'] + o['chegou'] - o['saiu'], 3)) for k, o in g.items()), key=lambda t: -t[2])
parado = [t for t in parado if t[2] > 0.01]

# ---------- planilha ----------
wb = Workbook()
B = lambda **k: Font(name=F, **k)  # noqa: E731
FILL = lambda c: PatternFill('solid', fgColor=c)  # noqa: E731
LINE = Border(bottom=Side(style='thin', color='E3E8EF'))


def white(ws, rows, cols):
    for r in range(1, rows):
        for c in range(1, cols):
            ws.cell(row=r, column=c).fill = FILL('FFFFFF')


def header_row(ws, row, heads, widths):
    for i, (h, w) in enumerate(zip(heads, widths), 1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = B(size=12, bold=True, color='FFFFFF')
        c.fill = FILL(NAVY)
        c.alignment = Alignment(vertical='center', wrap_text=True, horizontal='left')
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[row].height = 34


def back_link(ws, ref):
    c = ws[ref]
    c.value = '◄ Voltar ao início'
    c.font = B(size=11, bold=True, color=CORP, underline='single')
    c.hyperlink = "#'Comece Aqui'!A1"


# 1) COMECE AQUI
w1 = wb.active
w1.title = 'Comece Aqui'
w1.sheet_view.showGridLines = False
white(w1, 60, 10)
for col, w in zip('ABCDEFGH', [3, 30, 3, 30, 3, 30, 3, 3]):
    w1.column_dimensions[col].width = w
w1['B2'] = 'Estoque ZBR Importado'
w1['B2'].font = B(size=28, bold=True, color=NAVY)
w1['B3'] = 'Resumo fácil · movimentos de 9 de setembro a 1º de outubro de 2026'
w1['B3'].font = B(size=13, color=GREY)
w1.row_dimensions[2].height = 40
total, certos = len(linhas), sum(d['sit'] == '✔ Certo' for d in linhas)
urg = sum(d['prior'] == 'Urgente' for d in problemas)
kg_parado = sum(t[2] for t in parado if t[1]['um'] == 'kg')
CARDS = [('B', '✔ ESTÃO CERTOS', f'{certos} de {total}', f'{round(certos / total * 100)}% dos produtos e lotes', G_BG, G_TX),
         ('D', '⚠ TÊM DIFERENÇA', f'{len(problemas)}', 'veja a aba "Produtos"', Y_BG, Y_TX),
         ('F', '❗ URGENTES', f'{urg}', 'resolver primeiro (aba "O Que Fazer")', R_BG, R_TX)]
for col, lab, num, hint, bg, tx in CARDS:
    for r, (val, size, bold) in zip((5, 6, 7), ((lab, 12, True), (num, 34, True), (hint, 11, False))):
        c = w1[f'{col}{r}']
        c.value, c.fill = val, FILL(bg)
        c.font = B(size=size, bold=bold, color=tx)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
w1.row_dimensions[6].height = 48
w1.row_dimensions[7].height = 30
ton = f'{kg_parado / 1000:,.1f}'.replace(',', 'X').replace('.', ',').replace('X', '.')
w1['B9'] = (f'📦 No depósito 2008 (onde chega a importação) há {ton} toneladas paradas '
            f'em {len(parado)} lotes, esperando ir para o estoque de venda.')
w1['B9'].font = B(size=13, bold=True, color=Y_TX)
w1['B9'].fill = FILL(Y_BG)
w1.merge_cells('B9:F9')
w1.row_dimensions[9].height = 30
w1['B9'].alignment = Alignment(vertical='center', wrap_text=True, indent=1)

w1['B11'] = 'COMO LER'
w1['B11'].font = B(size=15, bold=True, color=CORP)
for i, t in enumerate(['Cada linha é um produto (código SKU) em um lote.',
                       '"Deveria ter" é o que o sistema diz que entrou menos o que saiu. "Tem" é o que está no estoque.',
                       'Se os dois números são iguais, está certo. Se não, existe uma diferença.',
                       'Cores: verde = certo · amarelo = sobrando · vermelho = faltando ou urgente.'], 12):
    w1[f'B{i}'] = '• ' + t
    w1[f'B{i}'].font = B(size=12, color=DARK)
    w1.merge_cells(f'B{i}:F{i}')

w1['B17'] = 'O QUE FAZER PRIMEIRO'
w1['B17'].font = B(size=15, bold=True, color=CORP)
r = 18
for d in [d for d in problemas if d['prior'] == 'Urgente'][:5]:
    w1[f'B{r}'] = f"❗ {d['sku']} (lote {d['lote']}): {d['aconteceu']} → {d['fazer']}"
    w1[f'B{r}'].font = B(size=12, color=R_TX)
    w1[f'B{r}'].alignment = Alignment(wrap_text=True, vertical='top')
    w1.merge_cells(f'B{r}:F{r}')
    w1.row_dimensions[r].height = 34
    r += 1
r += 1
w1[f'B{r}'] = 'ABAS DESTA PLANILHA (clique para abrir)'
w1[f'B{r}'].font = B(size=15, bold=True, color=CORP)
r += 1
for aba, txt in (('Produtos', 'Todos os produtos e lotes: certo, sobrando ou faltando, e o porquê.'),
                 ('O Que Fazer', 'Só os que têm problema, do mais urgente ao menos urgente.'),
                 ('Depósito 2008', 'Importação que chegou e ainda não foi para o estoque de venda.'),
                 ('Palavras', 'O que significa cada termo.')):
    c = w1[f'B{r}']
    c.value, c.hyperlink = f'→ {aba}', f"#'{aba}'!A1"
    c.font = B(size=13, bold=True, color=CORP, underline='single')
    w1[f'D{r}'] = txt
    w1[f'D{r}'].font = B(size=12, color=DARK)
    w1.merge_cells(f'D{r}:F{r}')
    r += 1
r += 1
w1[f'B{r}'] = ('Números calculados a partir do relatório completo "ZBR_Importado_Relatorio_Gestao.xlsx" '
               '(versão para analistas, com todas as contas).')
w1[f'B{r}'].font = B(size=10, italic=True, color=GREY)
w1.merge_cells(f'B{r}:F{r}')

# 2) PRODUTOS
w2 = wb.create_sheet('Produtos')
w2.sheet_view.showGridLines = False
w2['A1'] = 'Produtos e lotes'
w2['A1'].font = B(size=22, bold=True, color=NAVY)
w2['A2'] = 'Use as setas do cabeçalho para filtrar (por exemplo, só "▼ Faltando"). As linhas mais importantes estão em cima.'
w2['A2'].font = B(size=11, color=GREY)
back_link(w2, 'J1')
H2 = ['Situação', 'Produto (SKU)', 'Lote', 'Descrição', 'Deveria ter', 'Tem no estoque', 'Diferença', 'Unidade',
      'O que aconteceu', 'O que fazer', 'Prioridade']
header_row(w2, 4, H2, [14, 22, 13, 36, 13, 15, 12, 10, 60, 52, 12])
for i, d in enumerate(linhas, 5):
    vals = [d['sit'], d['sku'], d['lote'], d['desc'], d['deveria'], d['tem'], d['dif'], d['um'],
            d['aconteceu'], d['fazer'], d['prior']]
    for j, v in enumerate(vals, 1):
        c = w2.cell(row=i, column=j, value=v)
        c.font = B(size=11, color=DARK, bold=(j in (1, 2, 7)))
        c.border = LINE
        c.alignment = Alignment(vertical='center', wrap_text=(j in (9, 10)))
        if j in (5, 6):
            c.number_format = NUM
        if j == 7:
            c.number_format = SIGN
L2 = 4 + len(linhas)
w2.add_table(Table(displayName='tbProdutos', ref=f'A4:K{L2}', tableStyleInfo=TableStyleInfo(name='TableStyleLight1')))
for txt, bg, tx in (('✔ Certo', G_BG, G_TX), ('▲ Sobrando', Y_BG, Y_TX), ('▼ Faltando', R_BG, R_TX)):
    w2.conditional_formatting.add(f'A5:A{L2}', FormulaRule(formula=[f'$A5="{txt}"'], fill=FILL(bg), font=Font(color=tx, bold=True)))
w2.conditional_formatting.add(f'K5:K{L2}', FormulaRule(formula=['$K5="Urgente"'], fill=FILL(R_BG), font=Font(color=R_TX, bold=True)))
w2.conditional_formatting.add(f'K5:K{L2}', FormulaRule(formula=['$K5="Alta"'], fill=FILL(Y_BG), font=Font(color=Y_TX, bold=True)))
w2.conditional_formatting.add(f'A5:K{L2}', FormulaRule(formula=['MOD(ROW(),2)=0'], fill=FILL(ZEBRA)))
w2.freeze_panes = 'C5'

# 3) O QUE FAZER
w3 = wb.create_sheet('O Que Fazer')
w3.sheet_view.showGridLines = False
w3['A1'] = 'O que fazer'
w3['A1'].font = B(size=22, bold=True, color=NAVY)
w3['A2'] = 'Só os produtos com problema, do mais urgente ao menos urgente. Preencha "Quem vai fazer", "Até quando" e "Feito?".'
w3['A2'].font = B(size=11, color=GREY)
back_link(w3, 'I1')
H3 = ['Nº', 'Prioridade', 'Produto (SKU)', 'Lote', 'O problema', 'O que fazer', 'Quem vai fazer', 'Até quando', 'Feito?']
header_row(w3, 4, H3, [6, 12, 22, 13, 60, 52, 22, 13, 10])
for c in (7, 8, 9):
    w3.cell(row=4, column=c).fill = FILL(CORP)
for i, d in enumerate(problemas, 5):
    q = f"{abs(d['dif']):,.0f}".replace(',', '.')
    falta = 'Faltam' if d['dif'] < 0 else 'Sobram'
    vals = [i - 4, d['prior'], d['sku'], d['lote'], f"{falta} {q} {d['um']}. {d['aconteceu']}", d['fazer'], '', None, 'Não']
    for j, v in enumerate(vals, 1):
        c = w3.cell(row=i, column=j, value=v)
        c.font = B(size=11, color=DARK, bold=(j in (2, 3)))
        c.border = LINE
        c.alignment = Alignment(vertical='center', wrap_text=(j in (5, 6)))
        if j in (7, 8, 9):
            c.fill = FILL('EEF5FB')
    w3.cell(row=i, column=8).number_format = 'DD/MM/YYYY'
    w3.row_dimensions[i].height = 32
L3 = 4 + len(problemas)
w3.add_table(Table(displayName='tbOQueFazer', ref=f'A4:I{L3}', tableStyleInfo=TableStyleInfo(name='TableStyleLight1')))
dv = DataValidation(type='list', formula1='"Não,Sim"', allow_blank=False)
w3.add_data_validation(dv)
dv.add(f'I5:I{L3 + 200}')
w3.conditional_formatting.add(f'A5:I{L3}', FormulaRule(formula=['$I5="Sim"'], fill=FILL(G_BG), font=Font(color=G_TX)))
w3.conditional_formatting.add(f'B5:B{L3}', FormulaRule(formula=['$B5="Urgente"'], fill=FILL(R_BG), font=Font(color=R_TX, bold=True)))
w3.conditional_formatting.add(f'B5:B{L3}', FormulaRule(formula=['$B5="Alta"'], fill=FILL(Y_BG), font=Font(color=Y_TX, bold=True)))
w3[f'A{L3 + 2}'] = f'=COUNTIF(I5:I{L3},"Sim")&" de "&COUNTA(C5:C{L3})&" já resolvidos"'
w3[f'A{L3 + 2}'].font = B(size=13, bold=True, color=CORP)
w3.freeze_panes = 'D5'

# 4) DEPÓSITO 2008
w4 = wb.create_sheet('Depósito 2008')
w4.sheet_view.showGridLines = False
w4['A1'] = 'Depósito 2008: importação parada'
w4['A1'].font = B(size=22, bold=True, color=NAVY)
w4['A2'] = ('O depósito 2008 é onde a importação chega. Daqui ela deve ir para o estoque de venda (ZBR). '
            'Abaixo, o que chegou e ainda não foi.')
w4['A2'].font = B(size=11, color=GREY)
back_link(w4, 'H1')
H4 = ['Produto (SKU)', 'Lote', 'Descrição', 'Chegou', 'Já foi para venda', 'Ainda parado', 'Unidade', 'O que fazer']
header_row(w4, 4, H4, [22, 13, 36, 13, 16, 14, 10, 44])
for i, ((sku, lote), o, p) in enumerate(parado, 5):
    vals = [sku, lote, desc.get(sku, ''), o['ini'] + o['chegou'], o['saiu'], p, o['um'], 'Transferir para o estoque de venda (ZBR).']
    for j, v in enumerate(vals, 1):
        c = w4.cell(row=i, column=j, value=v)
        c.font = B(size=11, color=DARK, bold=(j in (1, 6)))
        c.border = LINE
        if j in (4, 5, 6):
            c.number_format = NUM
L4 = 4 + len(parado)
w4.add_table(Table(displayName='tbParado2008', ref=f'A4:H{L4}', tableStyleInfo=TableStyleInfo(name='TableStyleLight1')))
w4.conditional_formatting.add(f'A5:H{L4}', FormulaRule(formula=['MOD(ROW(),2)=0'], fill=FILL(ZEBRA)))
w4.conditional_formatting.add(f'F5:F{L4}', FormulaRule(formula=['$F5>0'], fill=FILL(Y_BG), font=Font(color=Y_TX, bold=True)))
w4[f'A{L4 + 2}'] = 'Total parado (kg):'
w4[f'A{L4 + 2}'].font = B(size=13, bold=True, color=CORP)
w4[f'F{L4 + 2}'] = f'=SUMIFS(F5:F{L4},G5:G{L4},"kg")'
w4[f'F{L4 + 2}'].number_format = NUM
w4[f'F{L4 + 2}'].font = B(size=13, bold=True, color=CORP)
w4.freeze_panes = 'C5'

# 5) PALAVRAS
w5 = wb.create_sheet('Palavras')
w5.sheet_view.showGridLines = False
w5['A1'] = 'O que significa cada palavra'
w5['A1'].font = B(size=22, bold=True, color=NAVY)
back_link(w5, 'C1')
header_row(w5, 3, ['Palavra', 'Significado'], [26, 100])
GLOS = [('SKU', 'Código do produto no sistema.'),
        ('Lote', 'Número que identifica um grupo do mesmo produto, fabricado junto.'),
        ('Estoque de venda (ZBR)', 'Onde fica a mercadoria pronta para vender e entregar.'),
        ('Depósito 2008', 'Onde a importação chega antes de ir para o estoque de venda.'),
        ('Deveria ter', 'O que entrou menos o que saiu, segundo o sistema (mais o estoque de 8/set, se informado).'),
        ('Tem no estoque', 'O que o sistema mostra guardado hoje nas posições do armazém.'),
        ('Diferença', '"Tem no estoque" menos "Deveria ter". Positivo = sobrando; negativo = faltando.'),
        ('Dar baixa', 'Registrar no sistema que a mercadoria saiu (foi vendida ou entregue).'),
        ('Doca', 'Área do armazém onde a mercadoria fica ao carregar ou descarregar.'),
        ('Estoque antigo', 'Mercadoria que já estava guardada antes de 9 de setembro.'),
        ('Urgente / Alta / Média / Baixa', 'Ordem de importância para resolver. Urgente = falta grande ou mercadoria não encontrada.'),
        ('kg / un / cx', 'Quilos, unidades e caixas.')]
for i, (a, b) in enumerate(GLOS, 4):
    w5[f'A{i}'], w5[f'B{i}'] = a, b
    w5[f'A{i}'].font = B(size=12, bold=True, color=NAVY)
    w5[f'B{i}'].font = B(size=12, color=DARK)
    for col in 'AB':
        w5[f'{col}{i}'].border = LINE
        if i % 2 == 0:
            w5[f'{col}{i}'].fill = FILL(ZEBRA)

for ws in wb.worksheets:
    ws.sheet_properties.tabColor = {'Comece Aqui': NAVY, 'Produtos': CORP, 'O Que Fazer': 'C00000',
                                    'Depósito 2008': 'FFC000', 'Palavras': GREY}[ws.title]
    ws.page_setup.orientation = 'portrait' if ws.title in ('Comece Aqui', 'Palavras') else 'landscape'
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.right.text = 'Página &P de &N'
wb.save(OUT)
print('OK', OUT, len(linhas), 'linhas,', len(problemas), 'problemas,', len(parado), 'lotes parados')
